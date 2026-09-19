import pytest
from openpyxl import load_workbook

from services.metrics_engine import store
from services.reporting.data import daily_report_data, final_report_data
from services.reporting.excel import export_workbook


@pytest.fixture
def conn(tmp_path):
    conn = store.connect(tmp_path / "events.db")
    store.insert_visit(conn, {
        "visit_id": "v1", "booth_id": "booth-01", "entered_at": "2026-09-19T10:00:00+00:00",
        "exited_at": "2026-09-19T10:00:08+00:00", "zone_path": "aisle→stand", "dwell_seconds": 8.0,
        "is_stopper": True, "is_staff": False, "gender_est": "female", "gender_conf": 0.9,
        "age_bracket": "18-35", "age_conf": 0.9, "position_trace": "10,10;20,20",
    })
    return conn


def test_daily_report_data_shapes(conn):
    data = daily_report_data(conn, "booth-01", "2026-09-19")
    assert data["summary"]["passersby"] == 1
    assert data["dwell_distribution"] == [8.0]
    assert data["hourly_traffic"][0]["hour"] == "10"


def test_excel_export_roundtrip(conn, tmp_path):
    out = export_workbook(conn, "booth-01", str(tmp_path / "export.xlsx"))
    wb = load_workbook(out)
    visits_ws = wb["visits"]
    assert visits_ws.max_row == 2  # header + 1 visit
    assert visits_ws["A2"].value == "v1"


def test_excel_export_empty_booth_has_headers_only(tmp_path):
    conn = store.connect(tmp_path / "empty.db")
    out = export_workbook(conn, "booth-01", str(tmp_path / "export.xlsx"))
    wb = load_workbook(out)
    assert wb["visits"].max_row == 1     # header row only
    assert wb["daily_summary"].max_row == 1


def test_excel_export_excludes_staff_visits(conn, tmp_path):
    store.insert_visit(conn, {
        "visit_id": "staff1", "booth_id": "booth-01", "entered_at": "2026-09-19T09:00:00+00:00",
        "exited_at": "2026-09-19T09:05:00+00:00", "zone_path": "table", "dwell_seconds": 300.0,
        "is_stopper": True, "is_staff": True, "gender_est": None, "gender_conf": None,
        "age_bracket": None, "age_conf": None, "position_trace": "",
    })
    out = export_workbook(conn, "booth-01", str(tmp_path / "export.xlsx"))
    wb = load_workbook(out)
    assert wb["visits"].max_row == 2  # still just the one non-staff visit from the fixture


def test_excel_export_includes_daily_summary_row(conn, tmp_path):
    daily_report_data(conn, "booth-01", "2026-09-19")  # upserts daily_summary as a side effect
    out = export_workbook(conn, "booth-01", str(tmp_path / "export.xlsx"))
    wb = load_workbook(out)
    summary_ws = wb["daily_summary"]
    assert summary_ws.max_row == 2
    assert summary_ws["A2"].value == "2026-09-19"
    assert summary_ws["B2"].value == 1  # passersby


def test_final_report_data_covers_each_requested_day(conn):
    store.insert_visit(conn, {
        "visit_id": "v2", "booth_id": "booth-01", "entered_at": "2026-09-20T11:00:00+00:00",
        "exited_at": "2026-09-20T11:00:03+00:00", "zone_path": "aisle", "dwell_seconds": 0.0,
        "is_stopper": False, "is_staff": False, "gender_est": None, "gender_conf": None,
        "age_bracket": None, "age_conf": None, "position_trace": "",
    })

    data = final_report_data(conn, "booth-01", ["2026-09-19", "2026-09-20"])

    assert len(data["days"]) == 2
    assert data["days"][0]["summary"]["passersby"] == 1
    assert data["days"][1]["summary"]["passersby"] == 1


def test_final_report_data_with_no_days_is_empty(conn):
    assert final_report_data(conn, "booth-01", []) == {"days": []}
